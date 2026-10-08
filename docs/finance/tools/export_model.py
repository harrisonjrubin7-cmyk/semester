"""Export the finance workbook's cells (formulas as text, inputs as values) for engine.js.

Usage as a module: export_model.export(path) -> {sheet: {"A1": value-or-"=formula"}}.
README, Approval_Matrix and Scenario_Results are left out: the first two are text, and Scenario_Results is a
pasted snapshot that the dashboard replaces with live runs. Check row 22 reads that snapshot, so the dashboard
ignores it (see SNAPSHOT_CHECK_ROW).
"""
import datetime
from openpyxl import load_workbook

SHEETS = ['Scenario_Control', 'Assumptions', 'Revenue', 'Costs', 'Headcount', 'Financials', 'Summary',
          'Unit_Economics', 'Board_Pack', 'Checks', 'Gate_Schedule']
SNAPSHOT_CHECK_ROW = 22
EPOCH = datetime.datetime(1899, 12, 30)


def export(path):
    wb = load_workbook(path)
    out = {}
    for name in SHEETS:
        cells = {}
        for row in wb[name].iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, datetime.datetime):
                    v = (v - EPOCH).days
                elif isinstance(v, datetime.date):
                    v = (datetime.datetime(v.year, v.month, v.day) - EPOCH).days
                cells[c.coordinate] = v
        out[name] = cells
    return out
