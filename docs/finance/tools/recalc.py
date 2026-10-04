"""Recalculate the finance workbook in LibreOffice for a chosen scenario and return a values-only workbook."""
import os, shutil, subprocess, tempfile
from openpyxl import load_workbook


def recalc(path, scenario=None, keep=None):
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, 'in.xlsx')
    if scenario is not None:
        wb = load_workbook(path)
        wb['Scenario_Control']['D3'] = scenario
        wb.save(src)
    else:
        shutil.copy(path, src)
    out = os.path.join(tmp, 'out')
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', out, src], check=True, capture_output=True)
    res = os.path.join(out, 'in.xlsx')
    if keep:
        shutil.copy(res, keep)
    return load_workbook(res, data_only=True)
