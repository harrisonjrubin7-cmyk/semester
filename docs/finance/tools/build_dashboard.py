"""Build docs/finance/dashboard.html: the workbook as an interactive model.

The page carries the workbook's own formulas (export_model.py) and a small spreadsheet engine (engine.js), so every
control in the page recalculates the real model in the browser. Nothing here edits the workbook.

Build-time work, all from the workbook:
  * the formulas and inputs, for the engine;
  * a catalogue of every changeable input (Assumptions register, scenario drivers, hiring plan) with its default,
    unit, section and basis, so the page's controls are generated from the workbook rather than written by hand;
  * the published results for each scenario (recalculated in LibreOffice), so the page can show "published" next to
    "your inputs";
  * a layout check that the cells the page reads still hold what it expects.

Usage: python3 build_dashboard.py [--fragment PATH]
  Writes ../dashboard.html (a standalone page). With --fragment, also writes the page body without the wrapper,
  which is what gets published as an artifact.  Prove the engine against the workbook with engine_parity.py.
"""
import datetime, json, os, re, subprocess, sys
sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import load_workbook
from recalc import recalc
from export_model import export

HERE = os.path.dirname(os.path.abspath(__file__))
MODEL = os.path.join(HERE, '..', 'semester-financial-model.xlsx')
TEMPLATE = os.path.join(HERE, 'dashboard_template.html')
ENGINE = os.path.join(HERE, 'engine.js')
OUT = os.path.join(HERE, '..', 'dashboard.html')

NAMES = ['Conservative', 'Base', 'Aggressive', 'Enterprise-delayed', 'High AI cost', 'Incident cost',
         'Gated plan', 'Go/no-go gates', 'Gates + gated hiring']
DESC = [
    'Fewer students, slower conversion, 40% fewer institutional logos at 85% of price, slower go-live, more churn.',
    'The staged plan of record: every driver at its planning value, no gates, hires on their base months.',
    'More registrations and conversion, 40% more logos at 110% of price, lower churn. An upside, not a target.',
    'Base economics with institutional signing six months later and go-live three months later than planned.',
    'Base with AI token usage 20% higher and vendor prices 15% higher: tests the allowance and gross margin.',
    'Base with a security incident: response cost, service credits and lost sales across the horizon.',
    'CFO-recommended staging of Base: flexible hires held until their gates, so burn follows evidence.',
    'Base economics with every revenue line held until its go/no-go gate month. No paid pilot before Year 2.',
    'Go/no-go gates and flexible hires released nine months after their base month. The lowest peak need.',
]
# Inputs shown as "key levers" at the top of the page: Assumptions IDs, and H<row> for a hire's base pay.
KEY = ['A-002', 'A-137', 'A-138', 'A-139', 'A-140', 'A-141', 'A-142', 'A-143', 'A-031', 'A-026', 'A-034', 'A-044',
       'A-011', 'A-012', 'A-160', 'A-162', 'A-163', 'A-164', 'H9', 'H10']
BASIS = {'FACT', 'PROPOSED', 'HYPOTHESIS', 'VENDOR', 'REVIEW', 'POLICY'}

# (sheet, cell, text the label must start with): the cells dashboard_template.html reads for results.
LAYOUT = [('Summary', 'A8', '1. Student'), ('Summary', 'A9', '2. Institutional'), ('Summary', 'A10', '3. Implementation'),
          ('Summary', 'A11', '4. AI usage'), ('Summary', 'A12', '5. Marketplace'), ('Summary', 'A13', '6. Career'),
          ('Summary', 'A14', '7. Alumni'), ('Summary', 'A15', 'Contra-revenue'), ('Summary', 'A16', 'Total revenue'),
          ('Summary', 'A23', 'Gross margin'), ('Summary', 'A29', 'EBITDA'), ('Summary', 'A40', 'ARR (recurring only)'),
          ('Summary', 'A44', 'Headcount'), ('Unit_Economics', 'A61', 'Burn multiple'), ('Unit_Economics', 'A64', 'Peak cumulative'),
          ('Board_Pack', 'A8', 'Revenue'), ('Board_Pack', 'A12', 'EBITDA'), ('Board_Pack', 'A15', 'Cash, end'),
          ('Board_Pack', 'A18', 'ARR, end'), ('Board_Pack', 'A25', 'Headcount'),
          ('Checks', 'A23', 'WARNING: cash stays above zero'), ('Checks', 'A24', 'WARNING: cash stays above the minimum'),
          ('Gate_Schedule', 'J5', '36-month people cost'), ('Gate_Schedule', 'A6', 'G1'), ('Gate_Schedule', 'A15', 'G10'),
          ('Scenario_Control', 'B10', 'Institutional new-logo volume'), ('Scenario_Control', 'B12', 'Enterprise signing delay'),
          ('Scenario_Control', 'B26', 'Slip added to every gate month'), ('Assumptions', 'B40', 'Pilot → annual agreement conversion')]


def check_layout(wb):
    for sheet, cell, start in LAYOUT:
        v = str(wb[sheet][cell].value)
        assert v.startswith(start), f'{sheet}!{cell} is {v!r}, the dashboard expects {start!r}: update LAYOUT and the template'


def catalogue(wb):
    """Every changeable input, from the workbook's own layout."""
    a = wb['Assumptions']; inputs = []; section = ''
    for r in range(5, a.max_row + 1):
        ida, label = a.cell(r, 1).value, a.cell(r, 2).value
        if isinstance(ida, str) and re.match(r'^\d+\. ', ida):
            section = ida
            continue
        if not (isinstance(ida, str) and re.match(r'^A-\d+$', ida)):
            continue
        cells, vals = [], []
        for c in (4, 5, 6):
            v = a.cell(r, c).value
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cells.append(f'{"DEF"[c - 4]}{r}'); vals.append(v)
        if not cells or ida == 'A-001':
            continue
        basis = next((a.cell(r, c).value for c in (5, 6, 7) if a.cell(r, c).value in BASIS), '')
        inputs.append(dict(id=ida, label=label, unit=a.cell(r, 3).value or '', section=section, basis=basis, cells=cells, vals=vals))
    sc = wb['Scenario_Control']; drivers = []
    for r in range(7, 27):
        drivers.append(dict(id=sc.cell(r, 1).value, row=r, label=sc.cell(r, 2).value, unit=sc.cell(r, 3).value or '',
                            vals=[sc.cell(r, c).value for c in range(4, 13)]))
    h = wb['Headcount']; hires = []
    for r in range(9, 53):
        hires.append(dict(row=r, role=h.cell(r, 1).value, fn=h.cell(r, 46).value, gate=h.cell(r, 48).value, flex=h.cell(r, 43).value,
                          qty=h.cell(r, 41).value, start=h.cell(r, 42).value, pay=h.cell(r, 45).value))
    return inputs, drivers, hires


def published(wb):
    s, u, ck = wb['Summary'], wb['Unit_Economics'], wb['Checks']
    return dict(rev3=s['E16'].value, arr3=s['E40'].value, peak=u['E64'].value, mincash=ck['B23'].value, lowmonths=ck['B24'].value)


def main():
    wbv = load_workbook(MODEL, data_only=True)
    check_layout(wbv)
    inputs, drivers, hires = catalogue(load_workbook(MODEL))
    assert all(d['id'] for d in drivers) and len(inputs) > 150, 'catalogue looks wrong'
    pub = []
    for k in range(1, 10):
        pub.append(published(recalc(MODEL, k)))
        print(f'scenario {k} read', file=sys.stderr)
    rev = subprocess.run(['git', '-C', HERE, 'rev-parse', '--short', 'HEAD'], capture_output=True, text=True).stdout.strip()
    data = dict(names=NAMES, desc=DESC, model=export(MODEL), inputs=inputs, drivers=drivers, hires=hires, key=KEY, pub=pub,
                stamp=f'{datetime.date.today().isoformat()} (workbook at {rev})')
    body = open(TEMPLATE).read().replace('/*DATA*/', json.dumps(data, separators=(',', ':'), default=str).replace('</', '<\\/'))
    body = body.replace('/*ENGINE*/', open(ENGINE).read())
    if '--fragment' in sys.argv:
        open(sys.argv[sys.argv.index('--fragment') + 1], 'w').write(body)
    title = body.split('<title>')[1].split('</title>')[0]
    page = ('<!doctype html><html lang="en"><head><meta charset="utf-8">'
            '<meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">'
            '</head><body>' + body + '</body></html>')
    open(OUT, 'w').write(page)
    print('wrote', os.path.normpath(OUT), f'{len(page):,}', 'bytes;', title, file=sys.stderr)


if __name__ == '__main__':
    main()
