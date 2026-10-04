"""Recalculate the workbook in LibreOffice and compare every model cell with the Python reference run.

Usage: python3 verify.py [workbook.xlsx] [--publish]   (default: ../semester-financial-model.xlsx)
--publish overwrites the workbook with the LibreOffice-recalculated copy (formulas kept, values cached) when verification passes.
Exit code 1 on any mismatch or failed in-workbook check.
"""
import json, os, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook
import engine, spec
from dsl import col

args = [a for a in sys.argv[1:] if not a.startswith('--')]
src = args[0] if args else os.path.join(os.path.dirname(__file__), '..', 'semester-financial-model.xlsx')
lay = json.load(open(os.path.join(os.path.dirname(__file__), 'layout.json')))
row_no, sheets = lay['row_no'], lay['sheets']

tmp = tempfile.mkdtemp()
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, src], check=True, capture_output=True)
rec = os.path.join(tmp, os.path.basename(src))
wb = load_workbook(rec, data_only=True)

bad, n = [], 0
for i, sname in enumerate(sheets):
    ws = wb[sname]
    out, kp = engine.run(i)
    for rd in spec.ROWS:
        if not rd.name: continue
        for t in range(1, 37):
            v = ws.cell(row_no[rd.name], 2 + t).value
            ref = out[rd.name][t - 1]
            n += 1
            if v is None or abs(float(v) - float(ref)) > 1e-6 * max(1, abs(ref)):
                bad.append((sname, rd.name, t, v, ref))
        if rd.kind in ('flow', 'stock'):
            for y in range(3):
                v = ws.cell(row_no[rd.name], 40 + y).value
                vals = out[rd.name][12 * y:12 * y + 12]
                ref = sum(vals) if rd.kind == 'flow' else vals[-1]
                n += 1
                if v is None or abs(float(v) - ref) > 1e-6 * max(1, abs(ref)):
                    bad.append((sname, rd.name + ' [annual]', y + 1, v, ref))
    for rd in spec.KPIS:
        for y in range(3):
            v = ws.cell(row_no[rd.name], 40 + y).value
            ref = kp[rd.name][y]
            n += 1
            if v is None or abs(float(v) - ref) > 1e-6 * max(1, abs(ref)):
                bad.append((sname, rd.name, y + 1, v, ref))

print(f'compared {n} cells across {len(sheets)} scenario sheets; mismatches: {len(bad)}')
for b in bad[:25]:
    print('  MISMATCH', b)

ck = wb['Checks']
fails = []
for r in range(4, ck.max_row + 1):
    for c in range(2, 2 + len(sheets)):
        v = ck.cell(r, c).value
        if v == 'FAIL' or (isinstance(v, str) and v.startswith('#')):
            fails.append((ck.cell(r, 1).value, ck.cell(3, c).value, v))
print('in-workbook checks failing:', len(fails))
for f in fails: print('  ', f)
# errors anywhere
errs = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('#') and c.value[1:4].isupper():
                errs += 1
                if errs < 10: print('  ERROR CELL', ws.title, c.coordinate, c.value)
print('error cells:', errs)
ok = not (bad or fails or errs)
if '--publish' in sys.argv:
    if not ok:
        print('NOT publishing: verification failed')
    else:
        shutil.copy(rec, src)   # the recalculated copy carries cached values and the same formulas
        print('published recalculated workbook over', src)
sys.exit(0 if ok else 1)
