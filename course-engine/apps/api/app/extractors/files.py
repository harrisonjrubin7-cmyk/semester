from __future__ import annotations

import csv
import shutil
import stat
import tempfile
import zipfile
from pathlib import Path, PurePosixPath

import pymupdf as fitz
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.core.config import settings
from app.extractors.base import ExtractedChunk
from app.services.storage import ALLOWED_EXTENSIONS


def extract_pdf(path: Path) -> list[ExtractedChunk]:
    chunks = []
    with fitz.open(path) as document:
        for index, page in enumerate(document, 1):
            text = page.get_text("text").strip()
            confidence = 1.0 if text else 0.0
            chunks.append(ExtractedChunk(text or "[OCR required]", "page", page_number=index, confidence=confidence))
    return chunks


def extract_docx(path: Path) -> list[ExtractedChunk]:
    doc = Document(path)
    chunks = [ExtractedChunk(p.text, "paragraph") for p in doc.paragraphs if p.text.strip()]
    for table_index, table in enumerate(doc.tables, 1):
        text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        chunks.append(ExtractedChunk(text, "table", metadata={"table": table_index}))
    return chunks


def extract_pptx(path: Path) -> list[ExtractedChunk]:
    deck = Presentation(path)
    chunks = []
    for number, slide in enumerate(deck.slides, 1):
        text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip())
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        chunks.append(ExtractedChunk("\n".join(filter(None, [text, notes])), "slide", slide_number=number))
    return chunks


def extract_xlsx(path: Path) -> list[ExtractedChunk]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    chunks = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        text = "\n".join(" | ".join("" if v is None else str(v) for v in row) for row in rows)
        chunks.append(ExtractedChunk(text, "table", sheet_name=sheet.title, cell_range=sheet.calculate_dimension()))
    return chunks


def extract_text(path: Path) -> list[ExtractedChunk]:
    return [ExtractedChunk(path.read_text(encoding="utf-8", errors="replace"), "document")]


def extract_csv(path: Path) -> list[ExtractedChunk]:
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        rows = list(csv.reader(handle))
    width = max((len(r) for r in rows), default=1)
    return [ExtractedChunk("\n".join(" | ".join(r) for r in rows), "table", sheet_name="CSV", cell_range=f"A1:{chr(64+min(width,26))}{len(rows)}")]


def safe_zip_members(path: Path, output_dir: Path) -> list[Path]:
    output_dir = output_dir.resolve()
    if output_dir.exists() and any(output_dir.iterdir()):
        raise ValueError("Extraction output directory must be empty")

    staging = Path(tempfile.mkdtemp(prefix=".archive-", dir=output_dir.parent))
    extracted_names: list[PurePosixPath] = []
    try:
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > settings.max_archive_entries:
                    raise ValueError(
                        f"Archive contains more than {settings.max_archive_entries} entries"
                    )

                seen: set[str] = set()
                expanded = 0
                for info in members:
                    member = PurePosixPath(info.filename)
                    canonical_name = member.as_posix()
                    if (
                        member.is_absolute()
                        or ".." in member.parts
                        or "\\" in info.filename
                        or canonical_name in {"", "."}
                    ):
                        raise ValueError("Unsafe archive path")
                    if canonical_name in seen:
                        raise ValueError(f"Archive contains duplicate member: {canonical_name}")
                    seen.add(canonical_name)
                    if stat.S_ISLNK((info.external_attr >> 16) & 0xFFFF):
                        raise ValueError(f"Archive contains symbolic link: {canonical_name}")
                    if info.file_size > settings.max_upload_bytes:
                        raise ValueError(
                            f"Archive member exceeds {settings.max_upload_bytes} bytes"
                        )
                    expanded += info.file_size
                    if expanded > settings.max_archive_expanded_bytes:
                        raise ValueError(
                            "Archive expanded data exceeds "
                            f"{settings.max_archive_expanded_bytes} bytes"
                        )

                actual_expanded = 0
                for info in members:
                    if info.is_dir():
                        continue
                    member = PurePosixPath(info.filename)
                    suffix = member.suffix.lower()
                    if suffix not in ALLOWED_EXTENSIONS or suffix == ".zip":
                        continue
                    target = staging.joinpath(*member.parts)
                    target.parent.mkdir(parents=True, exist_ok=True)
                    member_size = 0
                    with archive.open(info) as source, target.open("xb") as destination:
                        while block := source.read(1024 * 1024):
                            member_size += len(block)
                            actual_expanded += len(block)
                            if member_size > settings.max_upload_bytes:
                                raise ValueError(
                                    f"Archive member exceeds {settings.max_upload_bytes} bytes"
                                )
                            if actual_expanded > settings.max_archive_expanded_bytes:
                                raise ValueError(
                                    "Archive expanded data exceeds "
                                    f"{settings.max_archive_expanded_bytes} bytes"
                                )
                            destination.write(block)
                    extracted_names.append(member)
        except (zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
            raise ValueError("Invalid ZIP archive") from exc

        if output_dir.exists():
            output_dir.rmdir()
        staging.replace(output_dir)
        return [output_dir.joinpath(*member.parts) for member in extracted_names]
    except BaseException:
        shutil.rmtree(staging, ignore_errors=True)
        raise


def extractor_for(path: Path):
    return {
        ".pdf": extract_pdf, ".docx": extract_docx, ".pptx": extract_pptx,
        ".xlsx": extract_xlsx, ".csv": extract_csv, ".txt": extract_text, ".md": extract_text,
    }.get(path.suffix.lower())
