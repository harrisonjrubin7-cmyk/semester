from __future__ import annotations

from pathlib import Path

from docx import Document
from jinja2 import Environment, FileSystemLoader, select_autoescape
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas

TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "templates"


def export_guide_pdf(title: str, content: dict, citations: dict[str, dict], target: Path) -> Path:
    from weasyprint import HTML
    env = Environment(loader=FileSystemLoader(TEMPLATE_DIR), autoescape=select_autoescape(["html"]))
    html = env.get_template("study_guide.html.j2").render(title=title, content=content, citations=citations)
    target.parent.mkdir(parents=True, exist_ok=True)
    HTML(string=html, base_url=str(TEMPLATE_DIR)).write_pdf(target)
    return target


def export_guide_docx(title: str, content: dict, citations: dict[str, dict], target: Path) -> Path:
    document = Document()
    document.add_heading(title, 0)
    for section in content.get("sections", []):
        document.add_heading(section["heading"], 1)
        document.add_paragraph(section.get("explanation", ""))
        for point in section.get("key_points", []):
            document.add_paragraph(point, style="List Bullet")
        document.add_heading("Sources", 2)
        for citation_id in section.get("citation_ids", []):
            item = citations.get(str(citation_id), {})
            document.add_paragraph(item.get("label", str(citation_id)))
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def _position(index: int, reverse: bool = False) -> tuple[float, float]:
    width, height = letter
    margin, gutter = 36, 12
    card_w, card_h = (width - 2 * margin - gutter) / 2, (height - 2 * margin - 3 * gutter) / 4
    row, col = index // 2, index % 2
    if reverse:
        col = 1 - col
    return margin + col * (card_w + gutter), height - margin - (row + 1) * card_h - row * gutter


def _draw_card(canvas: Canvas, x: float, y: float, body: str, footer: str, front: bool) -> None:
    width, height = letter
    margin, gutter = 36, 12
    card_w, card_h = (width - 2 * margin - gutter) / 2, (height - 2 * margin - 3 * gutter) / 4
    canvas.setFillColor(HexColor("#111418") if front else HexColor("#F7F4EA"))
    canvas.setStrokeColor(HexColor("#C9A86A"))
    canvas.roundRect(x, y, card_w, card_h, 10, fill=1, stroke=1)
    canvas.setFillColor(HexColor("#F7F4EA") if front else HexColor("#17191D"))
    words, lines, current = body.split(), [], ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if stringWidth(candidate, "Helvetica", 10) < card_w - 28:
            current = candidate
        else:
            lines.append(current); current = word
    if current: lines.append(current)
    canvas.setFont("Helvetica", 10)
    text_y = y + card_h - 32
    for line in lines[:9]:
        canvas.drawString(x + 14, text_y, line); text_y -= 14
    canvas.setFont("Helvetica", 7)
    canvas.drawRightString(x + card_w - 14, y + 12, footer)


def export_flashcards_pdf(cards: list[dict], target: Path, flip: str = "long_edge") -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    canvas = Canvas(str(target), pagesize=letter)
    for start in range(0, len(cards), 8):
        batch = cards[start:start + 8]
        for index, card in enumerate(batch):
            _draw_card(canvas, *_position(index), card["front"], f"{card.get('id', start + index + 1)} · Q", True)
        canvas.showPage()
        for index, card in enumerate(batch):
            label = f"Difficulty {card.get('difficulty', 3)}/5 · {card.get('source_label', 'Source linked')}"
            _draw_card(canvas, *_position(index, flip == "long_edge"), card["back"], label, False)
        canvas.showPage()
    canvas.save()
    return target
